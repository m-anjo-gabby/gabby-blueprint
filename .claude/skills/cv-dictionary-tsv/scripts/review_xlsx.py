"""
コンテンツチーム向けの確認依頼Excel（登録データ一覧＋要確認＋未決の方針）を作る。
コンテンツチームは日本語話者・英語ネイティブの混在のため、日英併記にする
（「はじめに」は日本語・英語を別の列、確認事項は同じセルの上段に日本語・下段に英語。登録データ一覧は見出しだけ併記）。

使い方（リポジトリルートで実行。Python 3 + openpyxl が必要）:
  python .claude/skills/cv-dictionary-tsv/scripts/review_xlsx.py --dict <辞書TSV> [--out <出力xlsx>] [--title <表題>]
      [--title-en <英語の表題>] [--all-pending] [--ledger <台帳TSV>]

  --dict         生成済みの辞書TSV（CV辞書 一括登録用）
  --out          出力先（既定: 辞書TSVと同じフォルダの <辞書TSV名>_確認依頼.xlsx）
  --title        「はじめに」シートの表題に添える名前（既定: 辞書TSVのファイル名）
  --title-en     表題の英語（既定: --title と同じ）
  --all-pending  台帳の pending 行をすべて載せる（既定: 辞書TSVに含まれる語だけ）
  --ledger       台帳（既定: docs/cv-dictionary/review-ledger.tsv。見本の出力では sample/ の台帳を使う）

元データ:
  - 要確認: 台帳の pending 行（確認事項の英語は question_en。空の行は日本語だけになるため、出力前に Claude が書く）
  - 未決の方針: docs/cv-dictionary/open-policies.json（対象語が1つもない方針は載せない。英語は *_en）
見本: docs/cv-dictionary/sample/（サンプルデータで出力した確認依頼Excel）

補正結果サマリー（summary_xlsx.py）も、本ファイルの書式・日英併記の部品を使う。
"""
import argparse
import csv
import json
import math
import os
import re
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", ".."))
LEDGER = os.path.join(REPO, "docs", "cv-dictionary", "review-ledger.tsv")
POLICIES = os.path.join(REPO, "docs", "cv-dictionary", "open-policies.json")

# reference.md「7. cv_id」と同じ対応（表示名・IPA・例）
CV = [
    ("green_tea", "GREEN TEA", "/iː/ /i/", "meet, see, team"),
    ("silver_pin", "SILVER PIN", "/ɪ/", "sit, pink, build"),
    ("gray_day", "GRAY DAY", "/eɪ/", "day, name, tailor"),
    ("red_pepper", "RED PEPPER", "/ɛ/ /e/", "bed, send, professional"),
    ("black_cat", "BLACK CAT", "/æ/", "cat, happy, family"),
    ("purple_shirt", "PURPLE SHIRT", "/ɜːr/ /ɜr/ /ɝ/ /ər/", "bird, her, turn"),
    ("cup_of_mustard", "MUSTARD", "/ʌ/ /ə/", "up, mother, love"),
    ("olive_sock", "OLIVE SOCK", "/ɑː/ /ɑ/ /ɒ/ /ɑr/", "job, stop, father, car"),
    ("auburn_dog", "AUBURN DOG", "/ɔː/ /ɔ/", "dog, law, water"),
    ("turquoise_toy", "TURQUOISE TOY", "/ɔɪ/", "boy, oil, choice"),
    ("orange_door", "ORANGE DOOR", "/ɔːr/ /ɔr/", "door, for, born, course"),
    ("rose_boat", "ROSE BOAT", "/oʊ/ /əʊ/", "no, home, road"),
    ("wooden_hook", "WOODEN HOOK", "/ʊ/", "good, look, should"),
    ("blue_moon", "BLUE MOON", "/uː/ /u/", "you, room, truth"),
    ("brown_cow", "BROWN COW", "/aʊ/", "cow, house, now"),
    ("white_tie", "WHITE TIE", "/aɪ/", "by, night, like"),
]
CV_NAME = {c[0]: c[1] for c in CV}
POS_JA = {
    "NOUN": "名詞", "VERB": "動詞", "ADJ": "形容詞", "ADV": "副詞", "PREP": "前置詞",
    "PRON": "代名詞", "CONJ": "接続詞", "ART": "冠詞", "INT": "間投詞",
}
POS_EN = {"NOUN": "noun", "VERB": "verb", "ADJ": "adj.", "ADV": "adv.", "PREP": "prep.",
          "PRON": "pron.", "CONJ": "conj.", "ART": "article", "INT": "interj."}
CATEGORY_EN = {"R音化母音": "R-colored vowel", "発音の揺れ": "Pronunciation variant",
               "アクセント位置の揺れ": "Stress variant", "IPAとcv_idの不一致": "IPA / cv_id mismatch",
               "その他": "Other", "コンテンツチーム指摘": "Content team comment"}

# 回答欄の選択肢（回答の読み取りに使うため、文言を変えたら SKILL.md の反映手順も確認する）
ANS_FOLLOW = "方針の回答に従う / Follow the policy answer"
ANS_POLICY = [ANS_FOLLOW, "この語は現状のままでよい / Keep this word as is",
              "候補に変更 / Change to the candidate", "その他（コメント欄に記入） / Other (see comments)"]
ANS_WORD = ["現状のままでよい / Keep as is", "候補に変更 / Change to the candidate",
            "その他（コメント欄に記入） / Other (see comments)"]

FONT = "Meiryo UI"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color="0E3196")
f_sub = Font(name=FONT, size=11, bold=True, color="0E3196")
f_note = Font(name=FONT, size=9, color="666666")
fill_head = PatternFill("solid", fgColor="0E3196")
fill_input = PatternFill("solid", fgColor="FFF2CC")
fill_group = PatternFill("solid", fgColor="E8EDF8")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)

SHEET_POLICY = "①確認事項（方針） Policies"
SHEET_WORDS = "②確認事項（語別） Words"
SHEET_DICT = "③登録データ一覧 Data"
SHEET_CV = "Color Vowel"

# ------------------------------------------------------------
# 日英併記（同じセルの上段に日本語・下段に英語。間に薄い破線、英語は灰色）
# ------------------------------------------------------------
# 細い列（品詞など・幅10）でも折り返さない長さにする
BI_SEPARATOR = "- - - - -"
_in_sep = InlineFont(rFont=FONT, sz=8, color="BFBFBF")


def _inline(bold=False, color=None):
    return InlineFont(rFont=FONT, sz=10, b=bold or None, color=color)


def bi(ja, en, bold=False):
    """日英2段のセル値。片方しか無い・同じ文言なら1段"""
    ja = ja or ""
    en = en or ""
    if not en or en == ja:
        return ja or None
    if not ja:
        return en
    return CellRichText(TextBlock(_inline(bold), ja), TextBlock(_in_sep, f"\n{BI_SEPARATOR}\n"),
                        TextBlock(_inline(bold, "595959"), en))


def bi_head(ja, en):
    """見出し用の日英2行（見出しは白文字のため装飾しない）"""
    return ja if ja == en else f"{ja}\n{en}"


def _units(text):
    # 全角は半角2文字分として数える（列幅の単位に合わせる）
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in text)


def fit_height(ws, row, min_height=30):
    """折り返しを見込んで行の高さを決める（openpyxl は自動調整しないため概算する）"""
    tallest = 0
    for c in ws[row]:
        if c.value is None:
            continue
        width = ws.column_dimensions[c.column_letter].width or 10
        # 本文1行 = 14pt、日英の間の破線（8pt）= 10pt で見積もる
        h = sum(10 if seg == BI_SEPARATOR else 14 * max(1, math.ceil(_units(seg) / max(width - 1, 1)))
                for seg in str(c.value).split("\n"))
        tallest = max(tallest, h)
    ws.row_dimensions[row].height = max(min_height, tallest + 6)


# ------------------------------------------------------------
# 共通の書式
# ------------------------------------------------------------
def read_tsv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def cv_label(cv_id):
    return f"{CV_NAME.get(cv_id, cv_id)}（{cv_id}）"


def header(ws, row, cols, widths):
    for i, (name, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32


def body_cell(ws, row, col, value, align=wrap, input_cell=False, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = f_bold if bold else f_base
    c.alignment, c.border = align, border
    if input_cell:
        c.fill = fill_input
    return c


def label_cell(ws, row, value):
    body_cell(ws, row, 2, value, bold=True).fill = fill_group


def intro_sheet(ws, title_ja, title_en):
    """「はじめに」シートの枠（日本語・英語を別の列にする）。次に書く行番号を返す"""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 62
    ws.cell(row=2, column=2, value=title_ja).font = f_title
    ws.cell(row=3, column=2, value=title_en).font = f_sub
    label_cell(ws, 5, bi_head("項目", "Item"))
    body_cell(ws, 5, 3, "日本語", bold=True)
    body_cell(ws, 5, 4, "English", bold=True)
    return 6


def intro_rows(ws, r, rows, input_cell=False):
    """（見出し, 日本語, 英語）の行を書く。次に書く行番号を返す"""
    for k, ja, en in rows:
        label_cell(ws, r, k)
        body_cell(ws, r, 3, ja, input_cell=input_cell)
        body_cell(ws, r, 4, en, input_cell=input_cell)
        fit_height(ws, r, 20)
        r += 1
    return r


def intro_heading(ws, r, text):
    ws.cell(row=r, column=2, value=text).font = f_sub
    return r + 1


def dropdown(options):
    """選択肢のプルダウン。セルに付けたあと attach() でシートに登録する"""
    formula = '"' + ",".join(options) + '"'
    # Excel のリスト入力規則は255文字まで
    if len(formula) > 257:
        raise ValueError(f"選択肢が長すぎます（255文字まで）: {options}")
    return DataValidation(type="list", formula1=formula, allow_blank=True)


def attach(ws, *dvs):
    """セルに付けた入力規則だけをシートに登録する（対象セルの無い入力規則があると Excel が壊れたファイルとみなす）"""
    for dv in dvs:
        if dv.sqref.ranges:
            ws.add_data_validation(dv)


# ------------------------------------------------------------
# 方針の対象語
# ------------------------------------------------------------
def ledger_matches(row, m):
    if row["category"] != m.get("category", row["category"]):
        return False
    if "cvIds" in m and row["cv_id"] not in m["cvIds"]:
        return False
    q = row["question"]
    if not all(s in q for s in m.get("questionIncludes", [])):
        return False
    anys = m.get("questionIncludesAny")
    return not anys or any(s in q for s in anys)


def dict_matches(row, m):
    if "pos" in m and row["part_of_speech"] not in m["pos"]:
        return False
    return not ("wordPattern" in m and not re.fullmatch(m["wordPattern"], row["word_en"]))


def resolve_policies(policies, ledger, dict_rows):
    """方針ごとに対象語を集め、対象が1つもない方針は除く。台帳の行には方針ID（P-01…）を付ける。"""
    resolved = []
    for p in policies:
        m = p["match"]
        if m["source"] == "ledger":
            rows = [r for r in ledger if not r.get("policy") and ledger_matches(r, m)]
            if not rows:
                continue
            pid = f"P-{len(resolved) + 1:02d}"
            for r in rows:
                r["policy"] = pid
            counts = {}
            for r in rows:
                counts.setdefault(r["cv_id"], []).append(r["word_en"])
            main = max(counts, key=lambda k: len(counts[k]))
            current = cv_label(main) + "".join(
                f"\n※ {', '.join(sorted(set(ws)))} は {CV_NAME.get(cv, cv)}"
                for cv, ws in counts.items() if cv != main
            )
            current_en = cv_label(main).replace("（", " (").replace("）", ")") + "".join(
                f"\n* {', '.join(sorted(set(ws)))}: {CV_NAME.get(cv, cv)}"
                for cv, ws in counts.items() if cv != main
            )
            words = sorted({r["word_en"] for r in rows})
            count, count_en = f"{len(rows)}件", f"{len(rows)} words"
        else:
            words = sorted({r["word_en"] for r in dict_rows if dict_matches(r, m)})
            if not words:
                continue
            pid = f"P-{len(resolved) + 1:02d}"
            current, current_en = p["current"], p.get("current_en")
            count = f"{len(words)}{m.get('countLabel', '語')}"
            count_en = f"{len(words)} words"
        options = [f"{ja} / {en}" for ja, en in zip(p["options"], p.get("options_en", p["options"]))]
        resolved.append({**p, "id": pid, "words": ", ".join(words), "count": count, "count_en": count_en,
                         "current": current, "current_en": current_en, "options": options})
    return resolved


# ------------------------------------------------------------
# シート
# ------------------------------------------------------------
def build_intro(ws, title, title_en, dict_name, dict_count, policy_count, word_count):
    ws.title = "はじめに Introduction"
    r = intro_sheet(ws, f"ColorVowel辞書データ 確認のお願い（{title}）",
                    f"ColorVowel Dictionary: Review Request ({title_en})")
    r = intro_rows(ws, r, [
        (bi_head("対象データ", "Data"), dict_name, dict_name),
        (bi_head("登録件数", "Entries"), f"{dict_count}語（英単語＋品詞の組み合わせ単位）",
         f"{dict_count} entries (one per word + part of speech)"),
        (bi_head("確認していただきたいこと", "What to review"),
         f"① 方針の確認 {policy_count}件　② 語ごとの確認 {word_count}件　③ 登録データ全体の確認（気づいた点があれば）",
         f"① {policy_count} policies  ② {word_count} words  ③ All registered data (any comments are welcome)"),
    ])
    r = intro_rows(ws, r, [
        (bi_head("ご回答期限", "Due date"), None, None),
        (bi_head("ご回答者", "Reviewed by"), None, None),
    ], input_cell=True)
    r = intro_heading(ws, r + 1, "シートの構成と確認の進め方 / Sheets and how to review")
    r = intro_rows(ws, r, [
        (SHEET_POLICY,
         "複数の語にまとめて効く判断です。ここを先に決めていただくと、②の多くの語が自動的に決まります。"
         "「回答」列のプルダウンから選び、補足があれば「コメント」列に記入してください。",
         "Decisions that apply to many words at once. Deciding these first settles many words in ②. "
         "Choose from the drop-down in the \"Answer\" column and add notes in \"Comments\" if needed."),
        (SHEET_WORDS,
         "発音や Color Vowel の判断が分かれる語の一覧です。「方針」列に P-01 などがある語は①の回答に従うため、"
         "例外にしたい場合だけ記入してください。方針が空欄の語は、1語ずつ「回答」列を選んでください。"
         "変更する場合は「修正後」の列に正しい値を記入してください。",
         "Words whose pronunciation or Color Vowel is uncertain. Words with a policy ID (P-01, etc.) follow your answer "
         "in ①, so fill them in only for exceptions. For words without a policy, choose an answer for each word. "
         "If a value should change, write the correct value in the \"Corrected\" columns."),
        (SHEET_DICT,
         "登録済みの全データです。フィルターで絞り込めます。誤り・違和感があれば「指摘・コメント」列に記入してください"
         "（日本語訳の表現なども歓迎です）。「要確認ID」がある行は②に載っている語です。",
         "All registered data, with filters. Write any errors or concerns in the \"Comments\" column "
         "(comments on the Japanese translations are also welcome). Rows with a review ID are listed in ②."),
        (SHEET_CV, "Color Vowel と発音記号（IPA）の対応表です。判断の参考にしてください。",
         "Color Vowel to IPA reference table."),
    ])
    r = intro_heading(ws, r + 1, "記入のルール / How to fill in")
    c = ws.cell(row=r, column=2, value=bi_head("記入欄", "Input cells"))
    c.font, c.fill, c.border, c.alignment = f_bold, fill_input, border, wrap
    body_cell(ws, r, 3, "黄色のセルが記入欄です。それ以外のセルは変更しないでください（回答の取り込みに使います）。")
    body_cell(ws, r, 4, "Yellow cells are for your input. Please do not edit other cells (they are used to read your answers).")
    fit_height(ws, r, 20)
    r = intro_rows(ws, r + 1, [
        (bi_head("記入例", "Example"),
         "②で year を SILVER PIN に変えたい場合 → 回答「候補に変更」、修正後 Color Vowel「silver_pin」、"
         "修正後 発音記号は変更がなければ空欄のままで構いません。",
         "To change year to SILVER PIN in ②: Answer \"Change to the candidate\", Corrected Color Vowel \"silver_pin\". "
         "Leave Corrected IPA empty if it does not change."),
    ])
    r = intro_heading(ws, r + 1, "用語 / Terms")
    intro_rows(ws, r, [
        (bi_head("音節", "Syllables"), "ハイフンで区切った音のまとまり（例: ac-cel-er-ate）",
         "Sound units separated by hyphens (e.g. ac-cel-er-ate)"),
        (bi_head("第一アクセント音節", "Primary stress"), "一番強く読む音節が何番目か（例: ac-cel-er-ate → 2）",
         "Which syllable carries the primary stress (e.g. ac-cel-er-ate → 2)"),
        (bi_head("アクセント母音の綴り", "Stressed-vowel spelling"),
         "第一アクセント音節の中で母音を表している文字（例: cel → e）",
         "The letters spelling the vowel in the primary stressed syllable (e.g. cel → e)"),
        ("Color Vowel", "第一アクセント母音の Color Vowel（アプリで色として表示される）",
         "The Color Vowel of the primary stressed vowel (shown as a color in the app)"),
        (bi_head("発音記号", "IPA"), "米国英語の発音記号（IPA）。ˈ の直後が第一アクセント",
         "General American IPA. The syllable after ˈ carries the primary stress"),
        (bi_head("原形", "Lemma"), "語形変化した語の元の形（例: accelerated → accelerate）。原形そのものの語は空欄",
         "Base form of an inflected word (e.g. accelerated → accelerate). Empty for base forms"),
    ])


def build_policies(ws, policies):
    cols = [bi_head("ID", "ID"), bi_head("論点", "Topic"), bi_head("確認したいこと", "Question"),
            bi_head("対象の語", "Words"), bi_head("件数", "Count"), bi_head("現在の設定", "Current setting"),
            bi_head("候補", "Candidates"), bi_head("回答", "Answer"), bi_head("コメント", "Comments")]
    header(ws, 1, cols, [7, 22, 46, 36, 14, 28, 22, 28, 34])
    for i, p in enumerate(policies, start=2):
        vals = [p["id"], bi(p["title"], p.get("title_en"), bold=True), bi(p["question"], p.get("question_en")),
                p["words"], bi(p["count"], p["count_en"]), bi(p["current"], p.get("current_en")),
                bi(p["candidates"], p.get("candidates_en"))]
        for j, v in enumerate(vals, start=1):
            body_cell(ws, i, j, v, align=center if j in (1, 5) else wrap)
        body_cell(ws, i, 8, None, input_cell=True)
        body_cell(ws, i, 9, None, input_cell=True)
        dv = dropdown(p["options"])
        dv.add(ws.cell(row=i, column=8))
        attach(ws, dv)
        fit_height(ws, i, 75)
    ws.freeze_panes = "C2"
    n = len(policies) + 3
    ws.cell(row=n, column=1, value="※ 対象の語は今回のデータに含まれるもの。決まった方針は今後作成する辞書データにも適用します。").font = f_note
    ws.cell(row=n + 1, column=1, value="* Words listed are those in this data set. Decided policies also apply to future dictionary data.").font = f_note


def build_words(ws, ledger):
    cols = [bi_head("要確認ID", "Review ID"), bi_head("分類", "Category"), bi_head("方針", "Policy"),
            bi_head("英単語", "Word"), bi_head("品詞", "POS"), bi_head("音節", "Syllables"),
            bi_head("第一アクセント音節", "Primary stress"), bi_head("アクセント母音の綴り", "Stressed-vowel spelling"),
            "Color Vowel", bi_head("発音記号", "IPA"), bi_head("確認したいこと", "Question"), bi_head("回答", "Answer"),
            bi_head("修正後 Color Vowel", "Corrected Color Vowel"), bi_head("修正後 発音記号", "Corrected IPA"),
            bi_head("修正後 音節・アクセント位置", "Corrected syllables / stress"), bi_head("コメント", "Comments")]
    header(ws, 1, cols, [11, 15, 8, 13, 10, 15, 13, 14, 16, 18, 46, 26, 20, 18, 22, 30])
    ws.row_dimensions[1].height = 58
    ans_policy, ans_word = dropdown(ANS_POLICY), dropdown(ANS_WORD)
    cv_dv = DataValidation(type="list", allow_blank=True, formula1=f"='{SHEET_CV}'!$A$2:$A${len(CV) + 1}")
    ordered = sorted(ledger, key=lambda r: (r.get("policy") or "P-99", r["category"], r["word_en"]))
    for i, r in enumerate(ordered, start=2):
        pos = r["part_of_speech"]
        vals = [r["id"], bi(r["category"], CATEGORY_EN.get(r["category"])), r.get("policy"), r["word_en"],
                bi(POS_JA.get(pos, pos), POS_EN.get(pos, pos)), r["syllables"],
                int(r["primary_stress_syllable"]), r["stress_vowel_spelling"],
                f'{CV_NAME.get(r["cv_id"], r["cv_id"])}\n({r["cv_id"]})', r["phonetic_spelling"],
                bi(r["question"], r.get("question_en"))]
        for j, v in enumerate(vals, start=1):
            body_cell(ws, i, j, v, align=center if j in (1, 3, 5, 7, 8) else wrap, bold=j == 4)
        for j in range(12, 17):
            body_cell(ws, i, j, None, input_cell=True)
        (ans_policy if r.get("policy") else ans_word).add(ws.cell(row=i, column=12))
        cv_dv.add(ws.cell(row=i, column=13))
        if r.get("policy"):
            ws.cell(row=i, column=12).value = ANS_FOLLOW
        fit_height(ws, i)
    attach(ws, ans_policy, ans_word, cv_dv)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:P{len(ordered) + 1}"


def build_dict(ws, dict_rows, ledger_ids):
    # データ部は日本語のまま（英語の読み手は英単語・発音記号・Color Vowel で読める）。見出しだけ併記する
    cols = ["No", bi_head("英単語", "Word"), bi_head("品詞", "POS"), bi_head("日本語訳", "Japanese"),
            bi_head("音節", "Syllables"), bi_head("第一アクセント音節", "Primary stress"),
            bi_head("アクセント母音の綴り", "Stressed-vowel spelling"), "Color Vowel", "cv_id", bi_head("発音記号", "IPA"),
            bi_head("原形", "Lemma"), bi_head("要確認ID", "Review ID"), bi_head("指摘・コメント", "Comments")]
    header(ws, 1, cols, [6, 16, 10, 40, 18, 13, 14, 15, 15, 22, 13, 11, 36])
    ws.row_dimensions[1].height = 58
    rows = sorted(dict_rows, key=lambda r: (r["word_en"].lower(), r["part_of_speech"]))
    for i, r in enumerate(rows, start=2):
        vals = [i - 1, r["word_en"], POS_JA.get(r["part_of_speech"], r["part_of_speech"]), r["word_ja"],
                r["syllables"], int(r["primary_stress_syllable"]), r["stress_vowel_spelling"],
                CV_NAME.get(r["cv_id"], r["cv_id"]), r["cv_id"], r["phonetic_spelling"], r.get("lemma") or None,
                ledger_ids.get((r["word_en"], r["part_of_speech"]))]
        for j, v in enumerate(vals, start=1):
            body_cell(ws, i, j, v, align=center if j in (1, 3, 6, 7, 12) else wrap, bold=j == 2)
        body_cell(ws, i, 13, None, input_cell=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{len(rows) + 1}"


def build_cv(ws, count_sheet=None):
    cols = ["cv_id", "Color Vowel", bi_head("発音記号（IPA）", "IPA"), bi_head("例", "Examples")]
    widths = [18, 18, 26, 30]
    if count_sheet:
        cols.append(bi_head("今回のデータの件数", "Count in this data"))
        widths.append(14)
    header(ws, 1, cols, widths)
    for i, (cid, name, ipa, ex) in enumerate(CV, start=2):
        vals = [cid, name, ipa, ex] + ([f"=COUNTIF('{count_sheet}'!$I:$I,A{i})"] if count_sheet else [])
        for j, v in enumerate(vals, start=1):
            body_cell(ws, i, j, v, align=center if j == 5 else wrap, bold=j == 2)
    if count_sheet:
        total = len(CV) + 2
        body_cell(ws, total, 4, bi_head("合計", "Total"), bold=True)
        body_cell(ws, total, 5, f"=SUM(E2:E{total - 1})", align=center, bold=True)
    ws.freeze_panes = "A2"


def finish(wb, out):
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToHeight = 0
    wb.save(out)


def main():
    # Windows のコンソール（cp932）では IPA 記号を出力できないため UTF-8 に固定する
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--dict", required=True)
    ap.add_argument("--out")
    ap.add_argument("--title")
    ap.add_argument("--title-en")
    ap.add_argument("--all-pending", action="store_true")
    ap.add_argument("--ledger", default=LEDGER)
    args = ap.parse_args()

    dict_rows = read_tsv(args.dict)
    dict_keys = {(r["word_en"], r["part_of_speech"]) for r in dict_rows}
    ledger = [r for r in read_tsv(args.ledger) if r["status"] == "pending"
              and (args.all_pending or (r["word_en"], r["part_of_speech"]) in dict_keys)]
    with open(POLICIES, encoding="utf-8") as f:
        policies = resolve_policies(json.load(f)["policies"], ledger, dict_rows)

    stem = os.path.splitext(os.path.basename(args.dict))[0]
    out = args.out or os.path.join(os.path.dirname(os.path.abspath(args.dict)), f"{stem}_確認依頼.xlsx")
    title = args.title or stem

    wb = Workbook()
    build_intro(wb.active, title, args.title_en or title, os.path.basename(args.dict), len(dict_rows),
                len(policies), len(ledger))
    build_policies(wb.create_sheet(SHEET_POLICY), policies)
    build_words(wb.create_sheet(SHEET_WORDS), ledger)
    build_dict(wb.create_sheet(SHEET_DICT), dict_rows, {(r["word_en"], r["part_of_speech"]): r["id"] for r in ledger})
    build_cv(wb.create_sheet(SHEET_CV), SHEET_DICT)
    # 件数の数式（COUNTIF）を Excel で開いたときに計算させる
    wb.calculation.fullCalcOnLoad = True
    finish(wb, out)

    print(f"出力: {out}")
    print(f"登録データ {len(dict_rows)}行 / 方針 {len(policies)}件 / 語別 {len(ledger)}件"
          f"（うち方針に従う {sum(1 for r in ledger if r.get('policy'))}件）")
    for p in policies:
        print(f"  {p['id']} {p['title']}: {p['count']}")
    missing = [r["id"] for r in ledger if not r.get("question_en")]
    if missing:
        print(f"注意: 確認事項の英語（question_en）が空の行 {len(missing)}件（日本語だけで出力）: {', '.join(missing)}")
    missing_p = [p["id"] for p in policies if not p.get("question_en") or "options_en" not in p]
    if missing_p:
        print(f"注意: 方針の英語（question_en / options_en）が無い方針: {', '.join(missing_p)}")


if __name__ == "__main__":
    main()
