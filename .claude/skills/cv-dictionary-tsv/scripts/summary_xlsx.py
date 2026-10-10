"""
コンテンツチーム向けの補正結果サマリーExcel（ルールの変更＋語別の補正結果＋ご回答への対応＋ルール全体）を作る。
コンテンツチームは日本語話者・英語ネイティブの混在のため日英併記にする
（「はじめに」は日本語・英語を別の列、各シートの本文は同じセルの上段に日本語・下段に英語）。

使い方（リポジトリルートで実行。Python 3 + openpyxl が必要）:
  python .claude/skills/cv-dictionary-tsv/scripts/summary_xlsx.py --since <確定日> --notes <連絡事項JSON>
      --dict "<表示名>=<辞書TSV>" [--dict ...] --out <出力xlsx>

  --since  この日以降に確定した台帳の行・変更したルールを載せる（YYYY-MM-DD。回答を反映した日）
  --notes  docs/cv-dictionary/feedback/<確定日>.json（語ごとの補正以外の連絡事項）
  --dict   確認依頼のもとにした（補正前の）辞書TSV。「表示名=パス」で表示名を付けられる。複数指定可
  --out    出力先

元データ:
  - 語別の補正結果: docs/cv-dictionary/review-ledger.tsv の confirmed 行（補正前の値は --dict の辞書TSV）。
    decision_note の先頭の [R-xx] で、補正の根拠になったルールを示す
  - ルール: docs/cv-dictionary/rules.json（history の date が --since 以降の行を「今回の変更」とする）
  - 次回の確認予定: docs/cv-dictionary/open-policies.json
"""
import argparse
import json
import os
import re
import sys

from openpyxl import Workbook

# 同じフォルダの review_xlsx.py を部品として使う（リポジトリに __pycache__ を作らない）
sys.dont_write_bytecode = True
from review_xlsx import (  # noqa: E402
    CATEGORY_EN, CV_NAME, POLICIES, POS_EN, POS_JA, LEDGER, REPO,
    bi, bi_head, body_cell, build_cv, center, f_note, fill_input, finish, fit_height, header,
    intro_heading, intro_rows, intro_sheet, read_tsv, wrap,
)

RULES = os.path.join(REPO, "docs", "cv-dictionary", "rules.json")
VALUE_COLS = ["syllables", "primary_stress_syllable", "stress_vowel_spelling", "cv_id", "phonetic_spelling"]
SHEET_RULE_CHANGES = "①ルールの変更 Rule changes"
SHEET_WORDS = "②補正結果 Corrections"
SHEET_NOTES = "③ご回答への対応 Responses"
SHEET_RULES = "④作成ルール Rules"
SHEET_CV = "Color Vowel"
fill_changed = fill_input  # 変わった値・今回変更したルールは黄色で示す

STATUS_EN = {"確定": "Agreed", "開発側の基準": "Dev team default", "確認予定": "To be asked next"}
KIND_EN = {"確定": "Confirmed", "変更": "Changed", "追加": "Added"}
RULE_REF = re.compile(r"^\[(R-\d+)\]\s*")


def ordinal(n):
    n = int(n)
    return f"{n}{'st' if n == 1 else 'nd' if n == 2 else 'rd' if n == 3 else 'th'}"


def cv_text(cv_id):
    return f"{CV_NAME.get(cv_id, cv_id)}\n({cv_id})"


def stress_text(r):
    return bi(f'{r["syllables"]}\n{r["primary_stress_syllable"]}音節目（{r["stress_vowel_spelling"]}）',
              f'{ordinal(r["primary_stress_syllable"])} syllable ({r["stress_vowel_spelling"]})')


def change_en(b, a):
    """補正前後の差分を英語で書く（英語の読み手向けの「対応内容」）"""
    parts = []
    if b["cv_id"] != a["cv_id"]:
        parts.append(f'Color Vowel: {CV_NAME.get(b["cv_id"], b["cv_id"])} → {CV_NAME.get(a["cv_id"], a["cv_id"])}')
    if b["primary_stress_syllable"] != a["primary_stress_syllable"]:
        parts.append(f'Stress: {ordinal(b["primary_stress_syllable"])} → {ordinal(a["primary_stress_syllable"])} syllable')
    if b["syllables"] != a["syllables"] or b["stress_vowel_spelling"] != a["stress_vowel_spelling"]:
        parts.append(f'Syllables: {b["syllables"]} ({b["stress_vowel_spelling"]}) → {a["syllables"]} ({a["stress_vowel_spelling"]})')
    if b["phonetic_spelling"] != a["phonetic_spelling"]:
        parts.append(f'IPA: {b["phonetic_spelling"]} → {a["phonetic_spelling"]}')
    return "\n".join(parts) or "No change (confirmed as is)"


def split_note(note):
    """decision_note を（ルールID, 本文）に分ける。末尾の「（確認依頼 …）」は「対象データ」列と重なるため外す"""
    note = re.sub(r"（確認依頼 [^）]*）$", "", note)
    m = RULE_REF.match(note)
    return (m.group(1), note[m.end():]) if m else (None, note)


def changed_rules(rules, since):
    return [(r, h) for r in rules for h in r["history"] if h["date"] >= since]


def example(r):
    return bi(r["example"], r.get("example_en"))


def build_intro(ws, notes, dicts, words, rule_changes, next_policies, since):
    ws.title = "はじめに Introduction"
    title = notes.get("title") or " / ".join(dicts)
    r = intro_sheet(ws, f"ColorVowel辞書データ 補正結果のご報告（{title}）",
                    f"ColorVowel Dictionary: Correction Report ({notes.get('title_en') or title})")
    changed = sum(1 for w in words if w["changed"])
    per_dict = [(label, sum(1 for w in words if w["label"] == label and w["changed"])) for label in dicts]
    n_items = len(notes.get("items", []))
    requests = "\n".join(notes.get("requests", []))
    r = intro_rows(ws, r, [
        (bi_head("ご回答の反映日", "Applied on"), since, since),
        (bi_head("対象の確認依頼", "Review requests"), requests, requests),
        (bi_head("ルールの変更", "Rule changes"), f"{len(rule_changes)}件（変更・確定したルール）",
         f"{len(rule_changes)} rules changed or confirmed"),
        (bi_head("語ごとの補正", "Word corrections"),
         f"確定 {len(words)}語のうち、値を変えた語 {changed}語（" + "　".join(f"{lb} {n}語" for lb, n in per_dict) + "）",
         f"{changed} of {len(words)} confirmed words were changed (" + ", ".join(f"{lb}: {n}" for lb, n in per_dict) + ")"),
        (bi_head("ご回答への対応・保留", "Responses / open items"), f"{n_items}件（ご回答と異なる対応・解釈・ご相談を含みます）",
         f"{n_items} items, including where we handled your answer differently, how we interpreted it, and decisions we need from you"),
    ])
    r = intro_heading(ws, r + 1, "シートの構成 / Sheets")
    r = intro_rows(ws, r, [
        (SHEET_RULE_CHANGES,
         "今回のご回答で変わった・確定した辞書データ作成ルールです。今後作成する辞書データにも適用します。",
         "Dictionary rules that were changed or confirmed by your answers. They also apply to all future dictionary data."),
        (SHEET_WORDS,
         "ご回答を反映した語の一覧です（補正前 → 補正後）。黄色のセルが変わった値です。要確認に載せていない語も、"
         "「登録データ一覧」のご指摘やルールの変更で値を変えたものを載せています。",
         "Words updated from your answers (before → after). Yellow cells are changed values. Words that were not in the "
         "review list are included if your comments on the registered data or a rule change affected them."),
        (SHEET_NOTES,
         "ご回答と異なる対応をしたもの、ご回答をどう解釈したか、ご判断をお願いしたい保留事項です。特にご確認ください。"
         "ご意見があれば「ご意見」列に記入してお戻しください。",
         "Where we handled your answer differently, how we interpreted it, and items that need your decision. "
         "Please check this sheet in particular, and write any feedback in the \"Your comments\" column."),
        (SHEET_RULES, "現在の辞書データ作成ルールの全体です。「今回」列に印がある行が今回の変更です。",
         "All current dictionary rules. Rows marked in the \"This time\" column were changed this time."),
        (SHEET_CV, "Color Vowel と発音記号（IPA）の対応表です。", "Color Vowel to IPA reference table."),
    ])
    r = intro_heading(ws, r + 1, "次回の確認依頼でお尋ねする予定の論点 / Next review topics")
    r = intro_rows(ws, r, [(bi_head(p["title"], p.get("title_en") or p["title"]), p.get("question"), p.get("question_en"))
                           for p in next_policies] or [(bi_head("なし", "None"), None, None)])
    r = intro_heading(ws, r + 1, "アプリへの反映 / Applying to the app")
    intro_rows(ws, r, [(bi_head("音声", "Audio"),
                        "補正後の値は辞書データに上書き登録します。単語の音声は Azure の標準の読み上げ（英単語から自動作成）のため、"
                        "補正では作り直しません。発音を個別に調整した語だけ、音声を確認して作り直します。",
                        "Corrected values overwrite the dictionary data. Word audio is generated by Azure's standard TTS from the "
                        "word itself, so it is not regenerated by these corrections. Only words with individually tuned "
                        "pronunciation are checked and regenerated.")])


def build_rule_changes(ws, rule_changes):
    cols = [bi_head("ID", "ID"), bi_head("区分", "Area"), bi_head("ルール（現在）", "Rule (current)"),
            bi_head("例", "Examples"), bi_head("種類", "Type"), bi_head("変更前", "Before"), bi_head("経緯", "Background")]
    header(ws, 1, cols, [7, 16, 52, 40, 11, 32, 46])
    for i, (r, h) in enumerate(rule_changes, start=2):
        vals = [r["id"], bi(r["area"], r.get("area_en")), bi(r["rule"], r.get("rule_en"), bold=True), example(r),
                bi(h["kind"], KIND_EN.get(h["kind"])), bi(h["before"], h.get("before_en")), bi(h["note"], h.get("note_en"))]
        for j, v in enumerate(vals, start=1):
            c = body_cell(ws, i, j, v, align=center if j in (1, 5) else wrap)
            if j == 3 and h["kind"] != "確定":
                c.fill = fill_changed
        fit_height(ws, i)
    ws.freeze_panes = "C2"


def build_words(ws, words):
    cols = [bi_head("台帳ID", "Ledger ID"), bi_head("対象データ", "Data set"), bi_head("分類", "Category"),
            bi_head("英単語", "Word"), bi_head("品詞", "POS"), bi_head("結果", "Result"),
            bi_head("補正前 Color Vowel", "CV before"), bi_head("補正後 Color Vowel", "CV after"),
            bi_head("補正前 発音記号", "IPA before"), bi_head("補正後 発音記号", "IPA after"),
            bi_head("補正前 音節・アクセント", "Stress before"), bi_head("補正後 音節・アクセント", "Stress after"),
            bi_head("ルール", "Rule"), bi_head("対応内容", "Note")]
    header(ws, 1, cols, [10, 15, 15, 14, 10, 10, 17, 17, 17, 17, 20, 20, 8, 48])
    ws.row_dimensions[1].height = 45
    for i, w in enumerate(words, start=2):
        b, a = w["before"], w["after"]
        rule_id, note = split_note(a["decision_note"])
        pos = a["part_of_speech"]
        vals = [a["id"], w["label"], bi(a["category"], CATEGORY_EN.get(a["category"])), a["word_en"],
                bi(POS_JA.get(pos, pos), POS_EN.get(pos, pos)),
                bi("変更あり", "Changed") if w["changed"] else bi("変更なし", "No change")]
        for j, v in enumerate(vals, start=1):
            body_cell(ws, i, j, v, align=center if j in (1, 5, 6) else wrap, bold=j == 4)
        col = 7
        for before, after, diff in [
            (cv_text(b["cv_id"]), cv_text(a["cv_id"]), b["cv_id"] != a["cv_id"]),
            (b["phonetic_spelling"], a["phonetic_spelling"], b["phonetic_spelling"] != a["phonetic_spelling"]),
            (stress_text(b), stress_text(a), any(b[k] != a[k] for k in VALUE_COLS[:3])),
        ]:
            body_cell(ws, i, col, before)
            c = body_cell(ws, i, col + 1, after)
            if diff:
                c.fill = fill_changed
            col += 2
        body_cell(ws, i, 13, rule_id, align=center)
        reason = f"See rule {rule_id} (sheet ④)" if rule_id else "Per your answer / comment in the review request"
        body_cell(ws, i, 14, bi(note, f"{change_en(b, a)}\n{reason}"))
        fit_height(ws, i)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:N{len(words) + 1}"


def build_notes(ws, notes):
    cols = [bi_head("No", "No"), bi_head("種類", "Type"), bi_head("対象", "Subject"), bi_head("いただいたご回答", "Your answer"),
            bi_head("対応", "What we did"), bi_head("理由", "Why"), bi_head("ご意見", "Your comments")]
    header(ws, 1, cols, [5, 18, 24, 32, 46, 46, 32])
    for i, n in enumerate(notes.get("items", []), start=2):
        vals = [i - 1, bi(n["kind"], n.get("kind_en")), bi(n["target"], n.get("target_en"), bold=True), n["answer"],
                bi(n["action"], n.get("action_en")), bi(n["reason"], n.get("reason_en"))]
        for j, v in enumerate(vals, start=1):
            body_cell(ws, i, j, v, align=center if j == 1 else wrap)
        body_cell(ws, i, 7, None, input_cell=True)
        fit_height(ws, i)
    ws.freeze_panes = "D2"


def build_rules(ws, rules, since):
    cols = [bi_head("ID", "ID"), bi_head("区分", "Area"), bi_head("今回", "This time"), bi_head("ルール", "Rule"),
            bi_head("例", "Examples"), bi_head("状態", "Status"), bi_head("最終更新", "Updated")]
    header(ws, 1, cols, [7, 16, 9, 60, 46, 15, 12])
    for i, r in enumerate(rules, start=2):
        latest = max((h["date"] for h in r["history"]), default="")
        now = bool(latest) and latest >= since
        vals = [r["id"], bi(r["area"], r.get("area_en")), "●" if now else None, bi(r["rule"], r.get("rule_en"), bold=True),
                example(r), bi(r["status"], STATUS_EN.get(r["status"])), latest or None]
        for j, v in enumerate(vals, start=1):
            c = body_cell(ws, i, j, v, align=center if j in (1, 3, 6, 7) else wrap)
            if now and j in (3, 4):
                c.fill = fill_changed
        fit_height(ws, i)
    n = len(rules) + 3
    ws.cell(row=n, column=1, value="状態: 確定 = コンテンツチームと合意済み／開発側の基準 = 開発側で決めたもの（ご意見歓迎）"
                                   "／確認予定 = 次回の確認依頼でお尋ねするもの").font = f_note
    ws.cell(row=n + 1, column=1, value="Status: Agreed = agreed with the content team / Dev team default = set by the "
                                       "development team (feedback welcome) / To be asked next = will be asked in the next review request").font = f_note
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:G{len(rules) + 1}"


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--since", required=True)
    ap.add_argument("--notes", required=True)
    ap.add_argument("--dict", action="append", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    dicts = {}
    for spec in args.dict:
        label, _, path = spec.partition("=") if "=" in spec else ("", "", spec)
        label = label or os.path.splitext(os.path.basename(path))[0]
        dicts[label] = {(r["word_en"], r["part_of_speech"]): r for r in read_tsv(path)}

    words = []
    for a in read_tsv(LEDGER):
        if a["status"] != "confirmed" or a["decided_date"] < args.since:
            continue
        label = next((lb for lb, d in dicts.items() if (a["word_en"], a["part_of_speech"]) in d), None)
        if label is None:
            continue
        b = dicts[label][(a["word_en"], a["part_of_speech"])]
        words.append({"label": label, "before": b, "after": a,
                      "changed": any(b[k] != a[k] for k in VALUE_COLS)})
    words.sort(key=lambda w: (not w["changed"], w["after"]["category"], w["after"]["word_en"].lower()))

    with open(RULES, encoding="utf-8") as f:
        rules = json.load(f)["rules"]
    with open(args.notes, encoding="utf-8") as f:
        notes = json.load(f)
    with open(POLICIES, encoding="utf-8") as f:
        next_policies = json.load(f)["policies"]
    rule_changes = changed_rules(rules, args.since)

    wb = Workbook()
    build_intro(wb.active, notes, list(dicts), words, rule_changes, next_policies, args.since)
    build_rule_changes(wb.create_sheet(SHEET_RULE_CHANGES), rule_changes)
    build_words(wb.create_sheet(SHEET_WORDS), words)
    build_notes(wb.create_sheet(SHEET_NOTES), notes)
    build_rules(wb.create_sheet(SHEET_RULES), rules, args.since)
    build_cv(wb.create_sheet(SHEET_CV))
    finish(wb, args.out)

    print(f"出力: {args.out}")
    print(f"ルールの変更 {len(rule_changes)}件 / 語別 {len(words)}語（値を変えた語 {sum(w['changed'] for w in words)}語）"
          f" / 対応・保留 {len(notes.get('items', []))}件 / 次回の確認予定 {len(next_policies)}件")


if __name__ == "__main__":
    main()
