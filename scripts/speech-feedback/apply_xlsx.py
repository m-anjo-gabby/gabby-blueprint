"""
運営から返ってきた確認の Excel（review_xlsx.py で作ったもの）の「修正案」を feedbackMessages.json に取り込む。

使い方（リポジトリルートで実行。Python 3 + openpyxl が必要）:
  python scripts/speech-feedback/apply_xlsx.py --xlsx <回答のxlsx> [--dry-run]

- 修正案が空の行は変えない。
- 次の行は取り込まず、エラーとして一覧にする（1件でもあれば何も書き込まない）:
  - 差し込み語（{words} 等）の組が今の文と違う
  - 一覧を作った後に、その行の文がリポジトリ側で変わっている（「今の文」と JSON が違う）
  - ID が見つからない
- 「コメント」の書かれた行は、取り込みとは別に一覧にする（文の追加・削除や出し分けの相談は手で対応する）。
- 書き込んだ後に Prettier で JSON の書式をそろえる。
"""
import argparse
import os
import subprocess
import sys

from openpyxl import load_workbook

# 同じフォルダの catalog_rows を読み込む。リポジトリに __pycache__ を作らない
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(__file__))
from catalog_rows import CATALOG, REPO, get_at, load_catalog, placeholders, save_catalog, set_at  # noqa: E402
from review_xlsx import SHEET_LIST  # noqa: E402


def text_of(value):
    return "" if value is None else str(value).strip()


def read_answers(path):
    ws = load_workbook(path, read_only=True)[SHEET_LIST]
    rows = ws.iter_rows(values_only=True)
    header = [text_of(v) for v in next(rows)]
    col = {name: header.index(name) for name in ("ID", "今の文", "修正案", "コメント")}
    for values in rows:
        row_id = text_of(values[col["ID"]])
        if row_id:
            yield {name: text_of(values[i]) for name, i in col.items()}


def main():
    # Windows のコンソールの文字コードで日本語が化けないよう、出力を UTF-8 にそろえる
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True, help="運営から返ってきた xlsx")
    parser.add_argument("--dry-run", action="store_true", help="書き込まずに変更点だけを出す")
    args = parser.parse_args()

    data = load_catalog()
    changes, errors, comments = [], [], []
    for answer in read_answers(args.xlsx):
        row_id, proposal = answer["ID"], answer["修正案"]
        if answer["コメント"]:
            comments.append((row_id, answer["コメント"]))
        if not proposal:
            continue
        try:
            current = get_at(data, row_id)
        except (KeyError, IndexError, ValueError):
            errors.append((row_id, "ID が見つかりません"))
            continue
        if current != answer["今の文"]:
            errors.append((row_id, f"一覧の作成後に文が変わっています（今: {current}）"))
            continue
        if placeholders(proposal) != placeholders(current):
            errors.append(
                (row_id, f"差し込み語が今の文と違います（今: {placeholders(current)} / 修正案: {placeholders(proposal)}）")
            )
            continue
        if proposal != current:
            changes.append((row_id, current, proposal))

    for row_id, before, after in changes:
        print(f"[変更] {row_id}\n  前: {before}\n  後: {after}")
    for row_id, message in errors:
        print(f"[エラー] {row_id}: {message}")
    for row_id, comment in comments:
        print(f"[コメント] {row_id}: {comment}")
    print(f"変更 {len(changes)} 件 / エラー {len(errors)} 件 / コメント {len(comments)} 件")

    if errors:
        print("エラーがあるため書き込みませんでした。Excel の該当行を直すか、手で対応してから再実行してください。")
        sys.exit(1)
    if args.dry_run or not changes:
        return

    for row_id, _, after in changes:
        set_at(data, row_id, after)
    save_catalog(data)
    subprocess.run(["pnpm", "exec", "prettier", "--write", CATALOG], cwd=REPO, check=True, shell=os.name == "nt")
    print(f"{CATALOG} に書き込みました。単体テスト（pnpm --filter @gabby/testing unit）で確かめてください。")


if __name__ == "__main__":
    main()
