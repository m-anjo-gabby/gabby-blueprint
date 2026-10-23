"""
発話の評価の文言（feedbackMessages.json）を、運営に確認してもらう Excel の一覧にする。
運営は「修正案」「コメント」の列だけを書き、返ってきた Excel を apply_xlsx.py で取り込む。

使い方（リポジトリルートで実行。Python 3 + openpyxl が必要）:
  python scripts/speech-feedback/review_xlsx.py --out <出力xlsx>
"""
import argparse
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# 同じフォルダの catalog_rows を読み込む。リポジトリに __pycache__ を作らない
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(__file__))
from catalog_rows import TIER_LABELS, build_rows, load_catalog  # noqa: E402

SHEET_LIST = "コメント一覧"
# apply_xlsx.py はこの見出しの名前で列を探す
COLUMNS = [
    ("ID", 14),
    ("表示場所", 14),
    ("指摘の種類", 18),
    ("表示される場面", 40),
    ("今の文", 48),
    ("表示例", 56),
    ("修正案", 48),
    ("コメント", 32),
]

INTRO = [
    ("発話の評価コメント 確認のお願い", None),
    ("", None),
    ("単語帳・フレーズ・スプリントで発話したあとに出る評価のコメントの一覧です。", None),
    ("画面には、スコアの横の「一言」（1〜2文）と、その下の「改善のヒント」（最大2行）が出ます。", None),
    ("", None),
    ("書き方", "head"),
    ("・「コメント一覧」の「修正案」に、直したい文を書いてください。今のままでよい行は空欄のままにしてください。", None),
    ("・{words} {word} {heard} {pairs} {tip} {label} {confusionNote} は、単語などが入る場所です。消したり名前を変えたりせず、そのまま残してください（位置は変えてかまいません）。", None),
    ("・「表示例」は、差し込む場所に例の単語を入れた見え方です。一言の励まし（2文目）は、1文目の診断とつなげた形で示しています。", None),
    ("・文の追加・削除、出し分けの相談などは「コメント」に書いてください。", None),
    ("", None),
    ("コメントの決まり", "head"),
    ("・一言は「何がどう聞き取れたか（事実）」＋「その指摘に合った次の一歩」で書きます。指摘のあとに「この調子で」「そのまま」など、指摘を打ち消す言い方をしません。", None),
    ("・聞き取れなかった単語は「抜けている」「言えていない」と言い切らず、「聞き取れませんでした」と書きます（言っていても音声認識が拾わないことがあるため）。", None),
    ("・評価しているのは「単語が聞き取れたか」だけです。リズム・抑揚・流暢さを評価したように書きません。", None),
    ("・画面の見出しに合わせ、案内は「改善のヒント」と書きます。", None),
    ("・一言は狭い枠に出るため、1文を短くします（目安 40 字以内）。", None),
    ("", None),
    ("スコアの帯と一言の出し分け", "head"),
    (f"・{TIER_LABELS['excellent']}・{TIER_LABELS['great']}: 指摘は書かず、褒める文を1つ出します。", None),
    (f"・{TIER_LABELS['good']}・{TIER_LABELS['fair']}・{TIER_LABELS['poor']}: 最も優先する指摘1つの診断（1文目）＋その指摘と帯に合った励まし（2文目）を出します。", None),
    ("・指摘の優先順: 重要な単語 → 冠詞 → その他の単語。同じ語の中では「聞き取れない」→「別の語に聞こえた」→「2語に分かれた」。", None),
    ("・励ましに候補が複数ある場合は、毎回その中から1つを選んで出します。", None),
]

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="0E3196")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
WRAP = Alignment(wrap_text=True, vertical="top")


def write_intro(wb):
    ws = wb.active
    ws.title = "はじめに"
    ws.column_dimensions["A"].width = 120
    for i, (text, style) in enumerate(INTRO, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=14)
        elif style == "head":
            cell.font = Font(bold=True)
    ws.cell(row=len(INTRO) + 2, column=1, value=f"作成日: {date.today().isoformat()}")


def write_list(wb, rows):
    ws = wb.create_sheet(SHEET_LIST)
    for col, (name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = WRAP
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = width
    for r, row in enumerate(rows, 2):
        values = [row["id"], row["place"], row["group"], row["scene"], row["text"], row["example"], None, None]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.alignment = WRAP
            cell.border = BORDER
            if COLUMNS[col - 1][0] in ("修正案", "コメント"):
                cell.fill = INPUT_FILL
        ws.cell(row=r, column=1).font = Font(size=8, color="808080")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions


def main():
    # Windows のコンソールの文字コードで日本語が化けないよう、出力を UTF-8 にそろえる
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", required=True, help="出力する xlsx のパス")
    args = parser.parse_args()

    rows = build_rows(load_catalog())
    wb = Workbook()
    write_intro(wb)
    write_list(wb, rows)
    wb.save(args.out)
    print(f"{args.out}: {len(rows)} 行")


if __name__ == "__main__":
    main()
